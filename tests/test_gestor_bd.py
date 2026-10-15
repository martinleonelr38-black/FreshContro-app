# Pruebas del gestor de base de datos.
# Sirve para comprobar migraciones, reportes, stock vencido y reglas del inventario.
import os
import sqlite3
import sys
import tempfile
import unittest
from datetime import date
from openpyxl import load_workbook

APP_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "freshcontrol_app")
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from controladores.controlador_reportes import ControladorReportes
from controladores.controlador_inventario import ControladorInventario
from base_de_datos.gestor_bd import GestorBD
from base_de_datos import gestor_bd as gestor_bd_mod


class TestGestorBD(unittest.TestCase):
    def setUp(self):
        self.db_path = os.path.join(os.path.dirname(__file__), "test_freshcontrol.db")
        if os.path.exists(self.db_path):
            os.remove(self.db_path)

        self.gestor = GestorBD(db_path=self.db_path)
        conn = sqlite3.connect(self.db_path)
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE usuarios (
                id_usuario INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre TEXT NOT NULL,
                apellido TEXT NOT NULL,
                usuario TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                email TEXT NOT NULL,
                telefono TEXT NOT NULL,
                cargo TEXT,
                rol TEXT NOT NULL
            )
        """)
        cur.execute("""
            CREATE TABLE empleados (
                id_empleado INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre TEXT NOT NULL,
                apellido TEXT NOT NULL,
                usuario TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                email TEXT NOT NULL,
                telefono TEXT NOT NULL,
                cargo TEXT,
                rol TEXT NOT NULL
            )
        """)
        cur.execute("""
            CREATE TABLE proveedores (
                id_proveedor INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre TEXT NOT NULL,
                telefono TEXT,
                correo TEXT
            )
        """)
        cur.execute("""
            CREATE TABLE productos (
                id_producto INTEGER PRIMARY KEY AUTOINCREMENT,
                codigo_barras TEXT UNIQUE,
                nombre TEXT NOT NULL,
                categoria TEXT,
                precio REAL DEFAULT 0.0,
                cantidad_stock INTEGER DEFAULT 0,
                fecha_vencimiento DATE,
                id_proveedor INTEGER
            )
        """)
        cur.execute("""
            CREATE TABLE movimientos (
                id_movimiento INTEGER PRIMARY KEY AUTOINCREMENT,
                id_producto INTEGER NOT NULL,
                id_usuario INTEGER NOT NULL,
                tipo TEXT NOT NULL,
                cantidad INTEGER NOT NULL,
                fecha_hora DATETIME DEFAULT CURRENT_TIMESTAMP,
                detalle TEXT
            )
        """)
        cur.execute("INSERT INTO proveedores (nombre, telefono, correo) VALUES (?, ?, ?)", ("Proveedor 1", "123", "a@b.com"))
        cur.execute("INSERT INTO productos (codigo_barras, nombre, categoria, cantidad_stock, fecha_vencimiento, id_proveedor) VALUES (?, ?, ?, ?, ?, ?)", ("111", "Leche", "Lácteos", 20, "2026-10-10", 1))
        cur.execute("INSERT INTO movimientos (id_producto, id_usuario, tipo, cantidad, detalle) VALUES (?, ?, ?, ?, ?)", (1, 1, "entrada", 5, "Ingreso inicial"))
        conn.commit()
        conn.close()

    def tearDown(self):
        if os.path.exists(self.db_path):
            os.remove(self.db_path)

    def test_listar_productos_por_categoria(self):
        productos = self.gestor.listar_productos_por_categoria()
        self.assertTrue(len(productos) >= 1)
        self.assertEqual(productos[0][0], "Lácteos")

    def test_obtener_notificaciones_gerente(self):
        notifs = self.gestor.obtener_notificaciones_gerente()
        self.assertIsInstance(notifs, list)

    def test_registrar_proveedor(self):
        stock_inicial = self.gestor.ejecutar_consulta(
            "SELECT cantidad_stock FROM productos WHERE nombre = ?", ("Leche",)
        )[0][0]
        ok = self.gestor.registrar_proveedor("Proveedor 2", "456", "c@d.com")
        self.assertTrue(ok)
        proveedores = self.gestor.listar_proveedores()
        self.assertTrue(any(p[1] == "Proveedor 2" for p in proveedores))
        stock_final = self.gestor.ejecutar_consulta(
            "SELECT cantidad_stock FROM productos WHERE nombre = ?", ("Leche",)
        )[0][0]
        self.assertEqual(stock_final, stock_inicial)

    def test_reportes_no_incluyen_historial_de_proveedores(self):
        reporte = ControladorReportes()
        reporte.bd = self.gestor

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            reporte._ruta_reporte = lambda nombre: os.path.join(carpeta_temporal, nombre)
            ruta_excel = reporte.exportar_excel(
                tipo="movimientos", nombre_archivo="reporte_prueba.xlsx"
            )
            libro = load_workbook(ruta_excel, read_only=True)
            self.assertEqual(libro.sheetnames, ["Movimientos"])
            libro.close()

            ruta_pdf = reporte.exportar_pdf(
                tipo="movimientos", nombre_archivo="reporte_prueba.pdf"
            )
            self.assertTrue(os.path.isfile(ruta_pdf))
            self.assertGreater(os.path.getsize(ruta_pdf), 0)

    def test_reporte_pdf_incluye_vencidos_y_almacen_sin_historial_proveedor(self):
        self.gestor.ejecutar_comando(
            "INSERT INTO productos (codigo_barras, nombre, categoria, cantidad_stock, fecha_vencimiento, id_proveedor) VALUES (?, ?, ?, ?, ?, ?)",
            ("EXP-1", "Yogur vencido", "Lácteos", 3, "2020-01-01", 1),
        )
        reporte = ControladorReportes()
        reporte.bd = self.gestor

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            reporte._ruta_reporte = lambda nombre: os.path.join(carpeta_temporal, nombre)
            ruta_pdf = reporte.exportar_pdf(tipo="inventario", nombre_archivo="reporte_prueba.pdf")
            with open(ruta_pdf, "rb") as archivo:
                contenido = archivo.read()

            self.assertIn(b"Productos vencidos", contenido)
            self.assertIn(b"Inventario del almac", contenido)
            self.assertIn(b"Yogur vencido", contenido)
            self.assertNotIn(b"Historial de proveedores", contenido)

    def test_retirar_stock_vencido_conserva_producto_y_codigo(self):
        ayer = "2000-01-01"
        self.gestor.ejecutar_comando(
            "INSERT INTO productos (codigo_barras, nombre, categoria, cantidad_stock, fecha_vencimiento) VALUES (?, ?, ?, ?, ?)",
            ("expired-1", "Vencido uno", "General", 4, ayer),
        )
        id_primero = self.gestor.ejecutar_consulta(
            "SELECT id_producto FROM productos WHERE codigo_barras = ?", ("expired-1",)
        )[0][0]
        self.gestor.ejecutar_comando(
            "INSERT INTO movimientos (id_producto, id_usuario, tipo, cantidad, detalle) VALUES (?, ?, ?, ?, ?)",
            (id_primero, 1, "entrada", 4, "Ingreso anterior"),
        )
        self.gestor.ejecutar_comando(
            "INSERT INTO productos (codigo_barras, nombre, categoria, cantidad_stock, fecha_vencimiento) VALUES (?, ?, ?, ?, ?)",
            ("expired-2", "Vencido dos", "General", 2, ayer),
        )

        self.assertEqual(self.gestor.retirar_stock_vencido(id_primero), 1)
        producto_conservado = self.gestor.ejecutar_consulta(
            "SELECT codigo_barras, cantidad_stock, fecha_vencimiento FROM productos WHERE id_producto = ?",
            (id_primero,),
        )[0]
        self.assertEqual(producto_conservado, ("expired-1", 0, None))
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT COUNT(*) FROM movimientos WHERE id_producto = ?", (id_primero,)
            )[0][0],
            2,
        )
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT tipo, cantidad FROM movimientos WHERE id_producto = ? ORDER BY id_movimiento DESC LIMIT 1",
                (id_primero,),
            )[0],
            ("vencido", 4),
        )

        self.assertEqual(self.gestor.retirar_stock_vencido(), 1)
        self.assertEqual(self.gestor.listar_productos_vencidos(), [])
        self.gestor.registrar_movimiento(id_primero, 1, "entrada", 6, "Nuevo ingreso")
        producto_reabastecido = self.gestor.ejecutar_consulta(
            "SELECT codigo_barras, cantidad_stock FROM productos WHERE id_producto = ?", (id_primero,)
        )[0]
        self.assertEqual(producto_reabastecido, ("expired-1", 6))
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT COUNT(*) FROM productos WHERE codigo_barras = ?", ("expired-1",)
            )[0][0],
            1,
        )

    def test_quitar_fecha_conserva_producto_y_stock(self):
        self.assertTrue(self.gestor.eliminar_fecha_vencimiento(1))
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT nombre, cantidad_stock, fecha_vencimiento FROM productos WHERE id_producto = 1"
            )[0],
            ("Leche", 20, None),
        )

    def test_reporte_mensual_incluye_retiro_con_fecha_y_valor(self):
        fin = date.today().isoformat()
        self.gestor.ejecutar_comando(
            "ALTER TABLE movimientos ADD COLUMN fecha_vencimiento_producto DATE"
        )
        self.gestor.ejecutar_comando(
            "ALTER TABLE movimientos ADD COLUMN precio_unitario REAL DEFAULT 0"
        )
        self.gestor.ejecutar_comando(
            "INSERT INTO productos (codigo_barras, nombre, categoria, precio, cantidad_stock, fecha_vencimiento) VALUES (?, ?, ?, ?, ?, ?)",
            ("EXP-MONTH", "Yogur retirado", "Lácteos", 4.5, 0, fin),
        )
        producto_id = self.gestor.ejecutar_consulta(
            "SELECT id_producto FROM productos WHERE codigo_barras = ?", ("EXP-MONTH",)
        )[0][0]
        self.gestor.ejecutar_comando(
            "INSERT INTO movimientos (id_producto, id_usuario, tipo, cantidad, detalle, fecha_vencimiento_producto, precio_unitario) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (producto_id, 1, "vencido", 3, "Retiro por vencimiento", fin, 4.5),
        )
        reporte = ControladorReportes()
        reporte.bd = self.gestor
        resultado = reporte.obtener_vencidos_del_mes()
        yogur = next(fila for fila in resultado if fila[1] == "Yogur retirado")
        self.assertEqual(yogur[2:4], (3, 13.5))

        with tempfile.TemporaryDirectory() as carpeta_temporal:
            reporte._ruta_reporte = lambda nombre: os.path.join(carpeta_temporal, nombre)
            ruta = reporte.exportar_excel("vencidos_mensual", "vencidos.xlsx")
            libro = load_workbook(ruta, read_only=True)
            valores = list(libro["Vencidos del mes"].values)
            libro.close()
            self.assertIn(("Yogur retirado", 3, 13.5, fin), valores)
            ruta_pdf = reporte.exportar_pdf("vencidos_mensual", "vencidos.pdf")
            self.assertTrue(os.path.isfile(ruta_pdf))
            self.assertGreater(os.path.getsize(ruta_pdf), 0)

    def test_extiende_vencimiento_congelado_por_categoria(self):
        self.assertEqual(
            ControladorInventario.calcular_vencimiento_congelado("2024-08-31", "Carnes"),
            "2025-02-28",
        )

    def test_inicializacion_agrega_campos_de_ubicacion_congelado_y_snapshot(self):
        with tempfile.TemporaryDirectory() as carpeta_temporal:
            ruta_anterior = gestor_bd_mod.DB_PATH
            gestor_bd_mod.DB_PATH = os.path.join(carpeta_temporal, "migracion.db")
            try:
                gestor_bd_mod.inicializar_bd()
                gestor = GestorBD(db_path=gestor_bd_mod.DB_PATH)
                columnas_producto = {
                    fila[1] for fila in gestor.ejecutar_consulta("PRAGMA table_info(productos)")
                }
                columnas_movimiento = {
                    fila[1] for fila in gestor.ejecutar_consulta("PRAGMA table_info(movimientos)")
                }
                self.assertTrue({"pasillo", "imagen", "congelado", "fecha_congelacion"}.issubset(columnas_producto))
                self.assertTrue({"fecha_vencimiento_producto", "precio_unitario"}.issubset(columnas_movimiento))
            finally:
                gestor_bd_mod.DB_PATH = ruta_anterior

    def test_inicializacion_completa_esquema_sin_perder_datos(self):
        ruta_anterior = gestor_bd_mod.DB_PATH
        gestor_bd_mod.DB_PATH = self.db_path
        try:
            gestor_bd_mod.inicializar_bd()
        finally:
            gestor_bd_mod.DB_PATH = ruta_anterior

        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT nombre, telefono, correo FROM proveedores WHERE id_proveedor = 1"
            )[0],
            ("Proveedor 1", "123", "a@b.com"),
        )
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT codigo_barras, nombre, cantidad_stock FROM productos WHERE id_producto = 1"
            )[0],
            ("111", "Leche", 20),
        )
        tablas = {
            fila[0]
            for fila in self.gestor.ejecutar_consulta(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        self.assertTrue({"producto_proveedor", "alertas"}.issubset(tablas))
        self.assertNotIn("productos_legacy", tablas)
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT tipo, cantidad, detalle FROM movimientos WHERE id_movimiento = 1"
            )[0],
            ("entrada", 5, "Ingreso inicial"),
        )

    def test_alertas_persisten_se_listan_y_se_marcan_leidas(self):
        ruta_anterior = gestor_bd_mod.DB_PATH
        gestor_bd_mod.DB_PATH = self.db_path
        try:
            gestor_bd_mod.inicializar_bd()
        finally:
            gestor_bd_mod.DB_PATH = ruta_anterior

        self.gestor.sincronizar_alertas()
        self.gestor.sincronizar_alertas()
        alertas = self.gestor.listar_alertas()
        alerta_vencimiento = next(alerta for alerta in alertas if alerta[3] == "Vencimiento")
        self.assertEqual(alerta_vencimiento[2], "Leche")
        self.assertEqual(alerta_vencimiento[6], 0)
        self.assertEqual(
            self.gestor.ejecutar_consulta("SELECT COUNT(*) FROM alertas WHERE tipo_alerta = 'Vencimiento'")[0][0],
            1,
        )

        self.gestor.marcar_alerta_leida(alerta_vencimiento[0])
        alerta_actualizada = next(
            alerta for alerta in self.gestor.listar_alertas() if alerta[0] == alerta_vencimiento[0]
        )
        self.assertEqual(alerta_actualizada[6], 1)

    def test_stock_reducido_genera_una_notificacion_por_movimiento(self):
        ruta_anterior = gestor_bd_mod.DB_PATH
        gestor_bd_mod.DB_PATH = self.db_path
        try:
            gestor_bd_mod.inicializar_bd()
        finally:
            gestor_bd_mod.DB_PATH = ruta_anterior

        exito, mensaje = self.gestor.registrar_movimiento(
            1, 1, "salida", 2, "Prueba de salida de stock"
        )
        self.assertTrue(exito, mensaje)

        self.gestor.ejecutar_comando(
            """
            INSERT INTO productos
                (codigo_barras, nombre, categoria, cantidad_stock, fecha_vencimiento)
            VALUES (?, ?, ?, ?, ?)
            """,
            ("EXP-ALERTA", "Producto vencido de prueba", "General", 4, "2000-01-01"),
        )
        id_vencido = self.gestor.ejecutar_consulta(
            "SELECT id_producto FROM productos WHERE codigo_barras = ?",
            ("EXP-ALERTA",),
        )[0][0]
        self.assertEqual(self.gestor.retirar_stock_vencido(id_vencido, 1), 1)

        self.gestor.sincronizar_alertas()
        self.gestor.sincronizar_alertas()

        reducciones = [
            alerta for alerta in self.gestor.listar_alertas()
            if alerta[3] == "Stock reducido"
        ]
        self.assertEqual(len(reducciones), 2)
        mensajes = [alerta[4] for alerta in reducciones]
        self.assertTrue(any("salida" in mensaje.lower() for mensaje in mensajes))
        self.assertTrue(any("vencimiento" in mensaje.lower() for mensaje in mensajes))
        self.assertTrue(all("movimiento #" in mensaje for mensaje in mensajes))

    def test_inicializacion_elimina_legacy_y_conserva_movimientos(self):
        conn = sqlite3.connect(self.db_path)
        conn.execute(
            """
            CREATE TABLE productos_legacy (
                id_producto INTEGER PRIMARY KEY,
                codigo_barras TEXT,
                nombre TEXT NOT NULL,
                categoria TEXT,
                precio REAL,
                stock INTEGER,
                fecha_vencimiento DATE,
                id_proveedor INTEGER
            )
            """
        )
        conn.execute(
            """
            INSERT INTO productos_legacy
                (id_producto, codigo_barras, nombre, categoria, precio, stock, fecha_vencimiento, id_proveedor)
            VALUES (22, 'LEGACY-22', 'Producto legado', 'General', 8.5, 7, NULL, 1)
            """
        )
        conn.execute(
            """
            DROP TABLE movimientos
            """
        )
        conn.execute(
            """
            CREATE TABLE movimientos (
                id_movimiento INTEGER PRIMARY KEY AUTOINCREMENT,
                id_producto INTEGER NOT NULL,
                id_usuario INTEGER NOT NULL,
                tipo TEXT NOT NULL,
                cantidad INTEGER NOT NULL,
                fecha_hora DATETIME DEFAULT CURRENT_TIMESTAMP,
                detalle TEXT,
                FOREIGN KEY (id_producto) REFERENCES productos_legacy(id_producto)
            )
            """
        )
        conn.execute(
            """
            INSERT INTO movimientos (id_producto, id_usuario, tipo, cantidad, detalle)
            VALUES (22, 1, 'entrada', 7, 'Movimiento legado')
            """
        )
        conn.commit()
        conn.close()

        ruta_anterior = gestor_bd_mod.DB_PATH
        gestor_bd_mod.DB_PATH = self.db_path
        try:
            gestor_bd_mod.inicializar_bd()
        finally:
            gestor_bd_mod.DB_PATH = ruta_anterior

        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT codigo_barras, nombre, cantidad_stock FROM productos WHERE id_producto = 22"
            )[0],
            ("LEGACY-22", "Producto legado", 7),
        )
        self.assertEqual(
            self.gestor.ejecutar_consulta(
                "SELECT id_producto, tipo, cantidad, detalle FROM movimientos"
            )[0],
            (22, "entrada", 7, "Movimiento legado"),
        )
        tablas = {
            fila[0]
            for fila in self.gestor.ejecutar_consulta(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        self.assertNotIn("productos_legacy", tablas)
        claves_foraneas = self.gestor.ejecutar_consulta("PRAGMA foreign_key_list(movimientos)")
        self.assertEqual(claves_foraneas[0][2], "productos")


if __name__ == "__main__":
    unittest.main()
